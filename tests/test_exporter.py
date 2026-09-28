import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

# Ensure project root is in sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
import pytest

from src.core.exporter import export_to_excel, get_excel_bytes, COLUMNS_MAP
from src.config import OUTPUT_DIR, EXCEL_DIR


def test_export_to_excel_custom_path(tmp_path):
    sample_records = [
        {
            "reel_url": "https://www.facebook.com/reel/111",
            "caption": "Mô tả video 1 có link https://tintuc.vn/1",
            "found_in": "Mô tả",
            "target_url": "https://tintuc.vn/1",
            "title": "Tiêu đề tin tức 1",
            "content": "Nội dung chi tiết bài viết 1...",
            "video_path": "output/videos/111.mp4",
            "status": "Thành công",
            "scraped_at": "2026-09-28 08:30:00",
        }
    ]
    out_file = tmp_path / "test_result.xlsx"
    saved_path = export_to_excel(sample_records, out_file)
    assert Path(saved_path).exists()
    assert Path(saved_path) == out_file

    wb = openpyxl.load_workbook(saved_path)
    sheet = wb.active
    assert sheet.title == "Facebook Reels Content"
    assert sheet.freeze_panes == "A2"

    headers = [cell.value for cell in sheet[1]]
    expected_headers = list(COLUMNS_MAP.values())
    assert headers == expected_headers
    assert "Đường dẫn Video" in headers

    # Check header styling
    for col_idx in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col_idx)
        assert cell.font.bold is True
        assert cell.fill.start_color.rgb in ["001877F2", "1877F2"]

    # Check row data
    caption_col = headers.index("Tiêu đề Reels") + 1
    content_col = headers.index("Nội dung bài viết") + 1
    stt_col = headers.index("STT") + 1
    reel_url_col = headers.index("Link Reel") + 1
    video_col = headers.index("Đường dẫn Video") + 1

    assert sheet.cell(row=2, column=stt_col).value == 1
    assert sheet.cell(row=2, column=reel_url_col).value == "https://www.facebook.com/reel/111"
    assert sheet.cell(row=2, column=caption_col).value == "Mô tả video 1 có link https://tintuc.vn/1"
    assert sheet.cell(row=2, column=content_col).value == "Nội dung chi tiết bài viết 1..."
    assert sheet.cell(row=2, column=video_col).value == "output/videos/111.mp4"

    # Check alignment / wrap_text
    assert sheet.cell(row=2, column=content_col).alignment.wrap_text is True
    assert sheet.cell(row=2, column=stt_col).alignment.horizontal == "center"

    # Check column dimensions
    assert sheet.column_dimensions["A"].width == 8
    assert sheet.column_dimensions["C"].width == 40
    assert sheet.column_dimensions["F"].width == 65
    assert sheet.column_dimensions["G"].width == 35


def test_export_to_excel_default_path(tmp_path, monkeypatch):
    test_out_dir = tmp_path / "excel_out"
    monkeypatch.setattr("src.core.exporter.EXCEL_DIR", test_out_dir)

    sample_records = [
        {
            "reel_url": "https://www.facebook.com/reel/333",
            "title": "Default Path Test",
        }
    ]
    saved_path = export_to_excel(sample_records)
    assert Path(saved_path).exists()
    assert Path(saved_path).parent == test_out_dir
    assert Path(saved_path).name.startswith("reels_content_")
    assert Path(saved_path).suffix == ".xlsx"


def test_export_to_excel_empty_records(tmp_path):
    out_file = tmp_path / "empty.xlsx"
    saved_path = export_to_excel([], out_file)
    assert Path(saved_path).exists()

    wb = openpyxl.load_workbook(saved_path)
    sheet = wb.active
    assert sheet.max_row == 1
    headers = [cell.value for cell in sheet[1]]
    assert headers == list(COLUMNS_MAP.values())


def test_export_to_excel_missing_fields(tmp_path):
    sample_records = [
        {
            "reel_url": "https://www.facebook.com/reel/999",
        }
    ]
    out_file = tmp_path / "missing_fields.xlsx"
    saved_path = export_to_excel(sample_records, out_file)
    assert Path(saved_path).exists()

    wb = openpyxl.load_workbook(saved_path)
    sheet = wb.active
    assert sheet.cell(row=2, column=1).value == 1
    assert sheet.cell(row=2, column=2).value == "https://www.facebook.com/reel/999"
    assert sheet.cell(row=2, column=3).value in ("", None)
    # Scraped at is in the last column
    scraped_at_val = sheet.cell(row=2, column=len(COLUMNS_MAP)).value
    assert isinstance(scraped_at_val, str)
    assert len(scraped_at_val) > 0


def test_get_excel_bytes():
    sample_records = [{"reel_url": "https://www.facebook.com/reel/222", "title": "Bytes Test"}]
    data_bytes = get_excel_bytes(sample_records)
    assert isinstance(data_bytes, bytes)
    assert len(data_bytes) > 0

    wb = openpyxl.load_workbook(BytesIO(data_bytes))
    sheet = wb.active
    assert sheet.title == "Facebook Reels Content"
    headers = [cell.value for cell in sheet[1]]
    assert "Link Reel" in headers
    assert "Đường dẫn Video" in headers
    assert sheet.cell(row=2, column=headers.index("Link Reel") + 1).value == "https://www.facebook.com/reel/222"


def test_export_to_excel_with_output_dir(tmp_path):
    custom_dir = tmp_path / "custom_excel_folder"
    sample_records = [{"reel_url": "https://www.facebook.com/reel/444", "title": "Custom Dir Test"}]
    saved_path = export_to_excel(sample_records, output_dir=custom_dir)
    assert Path(saved_path).exists()
    assert Path(saved_path).parent == custom_dir
    assert Path(saved_path).name.startswith("reels_content_")
    assert Path(saved_path).suffix == ".xlsx"


def test_export_captions_to_txt(tmp_path):
    from src.core.exporter import export_captions_to_txt

    sample_records = [
        {"stt": 1, "caption": "Tiêu đề Reels 1 có link\nhttps://fb.com"},
        {"stt": 2, "caption": "Tiêu đề Reels 2 không dấu"},
        {"stt": 3, "caption": ""},
    ]
    txt_dir = tmp_path / "captions"
    count, saved_dir = export_captions_to_txt(sample_records, txt_dir)
    assert count == 3
    assert saved_dir == txt_dir
    assert (txt_dir / "caption_1.txt").exists()
    assert (txt_dir / "caption_2.txt").exists()
    assert (txt_dir / "caption_3.txt").exists()

    with open(txt_dir / "caption_1.txt", "r", encoding="utf-8") as f:
        assert f.read() == "Tiêu đề Reels 1 có link\nhttps://fb.com"
    with open(txt_dir / "caption_3.txt", "r", encoding="utf-8") as f:
        assert f.read() == ""



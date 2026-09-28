from datetime import datetime
from io import BytesIO
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.config import EXCEL_DIR, OUTPUT_DIR

COLUMNS_MAP = {
    "stt": "STT",
    "reel_url": "Link Reel",
    "caption": "Tiêu đề Reels",
    "found_in": "Vị trí tìm thấy link",
    "target_url": "Link Web",
    "content": "Nội dung bài viết",
    "video_path": "Đường dẫn Video",
    "status": "Trạng thái",
    "scraped_at": "Thời gian cào",
}


def _build_formatted_workbook(records: list[dict]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Facebook Reels Content"

    # 1. Ghi header
    headers = list(COLUMNS_MAP.values())
    ws.append(headers)

    # Style cho Header
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1877F2", end_color="1877F2", fill_type="solid")  # Facebook Blue
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    # 2. Ghi từng dòng dữ liệu
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    for idx, rec in enumerate(records, start=1):
        row_data = [
            idx,
            rec.get("reel_url", ""),
            rec.get("caption", ""),
            rec.get("found_in", ""),
            rec.get("target_url", ""),
            rec.get("content", "") or rec.get("title", ""),
            rec.get("video_path", ""),
            rec.get("status", ""),
            rec.get("scraped_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ]
        ws.append(row_data)
        current_row = idx + 1

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            header_name = headers[col_num - 1]
            if header_name in ["Tiêu đề Reels", "Nội dung bài viết"]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            elif header_name in ["STT", "Vị trí tìm thấy link", "Trạng thái", "Thời gian cào"]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(vertical="top")

    # 3. Điều chỉnh độ rộng cột
    column_widths = {
        "A": 8,   # STT
        "B": 35,  # Link Reel
        "C": 40,  # Tiêu đề Reels
        "D": 20,  # Vị trí tìm thấy link
        "E": 35,  # Link Web
        "F": 65,  # Nội dung bài viết
        "G": 35,  # Đường dẫn Video
        "H": 20,  # Trạng thái
        "I": 22,  # Thời gian cào
    }
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width

    return wb


def export_to_excel(
    records: list[dict],
    output_filepath: Path | str | None = None,
    output_dir: Path | str | None = None
) -> Path:
    """Lưu danh sách bản ghi ra file Excel (mặc định lưu tại EXCEL_DIR hoặc output_dir)."""
    if output_filepath is not None:
        p = Path(output_filepath)
        if p.is_dir() or (not p.suffix and not p.exists()):
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_filepath = p / f"reels_content_{timestamp}.xlsx"
        else:
            output_filepath = p
    else:
        target_dir = Path(output_dir) if output_dir else EXCEL_DIR
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_filepath = target_dir / f"reels_content_{timestamp}.xlsx"

    output_filepath.parent.mkdir(parents=True, exist_ok=True)
    wb = _build_formatted_workbook(records)
    wb.save(output_filepath)
    return output_filepath


def export_captions_to_txt(records: list[dict], target_dir: Path | str) -> tuple[int, Path]:
    """Lưu tiêu đề/caption của từng Reel thành các file caption_{số_thứ_tự}.txt trong target_dir.
    
    Returns:
        (count: int, saved_dir: Path)
    """
    target = Path(target_dir)
    target.mkdir(parents=True, exist_ok=True)
    count = 0
    for idx, rec in enumerate(records, start=1):
        stt = rec.get("stt") or idx
        caption = rec.get("caption", "") or ""
        txt_path = target / f"caption_{stt}.txt"
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(caption)
        count += 1
    return count, target


def get_excel_bytes(records: list[dict]) -> bytes:
    """Trả về buffer bytes Excel phục vụ tải về trực tiếp."""
    wb = _build_formatted_workbook(records)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
